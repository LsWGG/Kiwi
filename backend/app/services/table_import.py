"""表格文件导入：CSV / Excel → 数据表笔记的字段与记录。

CSV 使用标准库解析并自动识别常用分隔符与编码；Excel 通过 openpyxl / xlrd
读取。导入前统一做文件大小、压缩包展开上限、行列数与单元格长度约束。
"""

from __future__ import annotations

import csv
import io
import re
import uuid
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from app.services.file_security import scan_bytes

TABLE_EXT = {".csv", ".xlsx", ".xls"}
MAX_TABLE_BYTES = 20 * 1024 * 1024
MAX_TABLE_ROWS = 5000
MAX_TABLE_COLUMNS = 120
MAX_CELL_CHARS = 20000
MAX_SAMPLE_ROWS = 8
MAX_XLSX_UNCOMPRESSED = 200 * 1024 * 1024
MAX_XLSX_ENTRIES = 2000

_NUMBER_RE = re.compile(r"^[+-]?(?:\d+(?:\.\d*)?|\.\d+)$")
_URL_RE = re.compile(r"^https?://", re.IGNORECASE)
_TRUE_VALUES = {"true", "yes", "y", "是", "✓", "√"}
_FALSE_VALUES = {"false", "no", "n", "否", "×", "✗"}
_DATE_FORMATS = (
    "%Y-%m-%d",
    "%Y/%m/%d",
    "%Y.%m.%d",
    "%Y年%m月%d日",
    "%m/%d/%Y",
    "%d/%m/%Y",
)


class TableImportError(Exception):
    def __init__(self, message: str, status_code: int = 422):
        super().__init__(message)
        self.message = message
        self.status_code = status_code


@dataclass
class SheetInfo:
    name: str
    row_count: int
    column_count: int


@dataclass
class ParsedTable:
    sheets: list[SheetInfo]
    sheet: str
    fields: list[dict]
    rows: list[dict]
    truncated: bool


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()[:MAX_CELL_CHARS]


def _is_empty(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _decode_csv(content: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030", "utf-16", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise TableImportError("无法识别 CSV 编码，请另存为 UTF-8 后重试", 415)


def _read_csv(content: bytes) -> tuple[list[str], list[list[Any]]]:
    text = _decode_csv(content)
    if not text.strip():
        raise TableImportError("CSV 文件为空", 400)
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    grid: list[list[Any]] = []
    try:
        for index, row in enumerate(reader):
            if index > MAX_TABLE_ROWS + 1:
                break
            grid.append([str(cell)[:MAX_CELL_CHARS] for cell in row])
    except csv.Error as exc:
        raise TableImportError(f"CSV 解析失败：{exc}", 422) from exc
    return ["CSV"], grid


def _guard_xlsx(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as zf:
            infos = zf.infolist()
    except zipfile.BadZipFile as exc:
        raise TableImportError("Excel 文件已损坏或不是有效的 .xlsx", 415) from exc
    if len(infos) > MAX_XLSX_ENTRIES:
        raise TableImportError("Excel 内部条目过多，已拒绝解析", 422)
    if sum(info.file_size for info in infos) > MAX_XLSX_UNCOMPRESSED:
        raise TableImportError("Excel 解压后体积超过 200MB 上限", 422)


def _read_xlsx(content: bytes) -> tuple[list[str], dict[str, list[list[Any]]]]:
    _guard_xlsx(content)
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - 依赖缺失
        raise TableImportError("服务端缺少 Excel 解析组件 openpyxl", 500) from exc

    try:
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=True
        )
    except Exception as exc:  # noqa: BLE001 - 第三方解析异常统一转 415
        raise TableImportError("Excel 文件解析失败，请确认文件未加密", 415) from exc

    names: list[str] = []
    sheets: dict[str, list[list[Any]]] = {}
    try:
        for worksheet in workbook.worksheets:
            names.append(worksheet.title)
            grid: list[list[Any]] = []
            for row_index, row in enumerate(worksheet.iter_rows(values_only=True)):
                if row_index > MAX_TABLE_ROWS + 1:
                    break
                grid.append([_cell_text(cell) for cell in row[:MAX_TABLE_COLUMNS]])
            sheets[worksheet.title] = grid
    finally:
        workbook.close()
    return names, sheets


def _read_xls(content: bytes) -> tuple[list[str], dict[str, list[list[Any]]]]:
    try:
        import xlrd
    except ImportError as exc:  # pragma: no cover - 依赖缺失
        raise TableImportError("服务端缺少旧版 Excel 解析组件 xlrd", 500) from exc
    try:
        workbook = xlrd.open_workbook(file_contents=content)
    except Exception as exc:  # noqa: BLE001
        raise TableImportError("Excel 文件解析失败，请确认文件未加密", 415) from exc

    names = workbook.sheet_names()
    sheets: dict[str, list[list[Any]]] = {}
    for name in names:
        sheet = workbook.sheet_by_name(name)
        grid: list[list[Any]] = []
        for row_index in range(min(sheet.nrows, MAX_TABLE_ROWS + 2)):
            row = []
            for col_index in range(min(sheet.ncols, MAX_TABLE_COLUMNS)):
                cell = sheet.cell(row_index, col_index)
                value: Any = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        value = xlrd.xldate.xldate_as_datetime(value, workbook.datemode)
                    except Exception:  # noqa: BLE001
                        pass
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    value = bool(value)
                row.append(_cell_text(value))
            grid.append(row)
        sheets[name] = grid
    return names, sheets


def _read_workbook(content: bytes, filename: str) -> tuple[list[str], dict[str, list[list[Any]]]]:
    ext = Path(filename).suffix.lower()
    if ext == ".csv":
        names, grid = _read_csv(content)
        return names, {"CSV": grid}
    if ext == ".xlsx":
        return _read_xlsx(content)
    if ext == ".xls":
        return _read_xls(content)
    raise TableImportError(
        f"不支持的表格类型 {ext or '（无扩展名）'}，允许：.csv、.xlsx、.xls", 415
    )


def _to_number(value: Any) -> int | float | None:
    if isinstance(value, bool) or _is_empty(value):
        return None
    if isinstance(value, (int, float)):
        return value
    text = str(value).strip()
    if not _NUMBER_RE.match(text):
        return None
    number = float(text)
    return int(number) if number.is_integer() else number


def _to_date(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    if not text:
        return None
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    # Excel 常把日期单元格读成 "2025-02-10 00:00:00" 这类完整时间串。
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).date().isoformat()
    except ValueError:
        return None


def _infer(field_id: str, name: str, values: list[Any]) -> dict:
    present = [value for value in values if not _is_empty(value)]
    if not present:
        return {"id": field_id, "name": name, "type": "text"}

    numbers = [_to_number(value) for value in present]
    if all(number is not None for number in numbers):
        lowered = name.lower()
        is_progress_header = any(
            marker in lowered for marker in ("进度", "完成度", "percent", "progress")
        )
        if is_progress_header and all(0 <= number <= 100 for number in numbers):
            return {"id": field_id, "name": name, "type": "progress"}
        return {"id": field_id, "name": name, "type": "number"}
    if all(_to_date(value) is not None for value in present):
        return {"id": field_id, "name": name, "type": "date"}
    if all(str(value).strip().lower() in _TRUE_VALUES | _FALSE_VALUES for value in present):
        return {"id": field_id, "name": name, "type": "checkbox"}
    if all(_URL_RE.match(str(value).strip()) for value in present):
        return {"id": field_id, "name": name, "type": "url"}
    if any(len(str(value)) > 200 for value in present):
        return {"id": field_id, "name": name, "type": "long_text"}
    return {"id": field_id, "name": name, "type": "text"}


def _typed_value(field: dict, value: Any) -> Any:
    ftype = field["type"]
    if ftype == "number":
        return _to_number(value)
    if ftype == "progress":
        number = _to_number(value)
        if number is None:
            return None
        return max(0, min(100, int(round(number))))
    if ftype == "date":
        return _to_date(value)
    if ftype == "checkbox":
        text = str(value).strip().lower()
        return text in _TRUE_VALUES if text else False
    text = _cell_text(value)
    return text


def _normalize_headers(raw: list[Any], width: int) -> list[str]:
    headers: list[str] = []
    seen: dict[str, int] = {}
    for index in range(width):
        value = _cell_text(raw[index]) if index < len(raw) else ""
        name = re.sub(r"\s+", " ", value).strip()[:120] or f"列{index + 1}"
        count = seen.get(name, 0) + 1
        seen[name] = count
        headers.append(name if count == 1 else f"{name} ({count})")
    return headers


# ---------- Excel 工作簿（多工作表 + 公式 + 基础样式） ----------
# 新的「表格」笔记类型直接保存工作簿 JSON；与数据表导入共用同一套
# 文件大小 / 行列数 / 单元格长度约束。前端按 A1 记法渲染与计算公式。

MAX_SHEETS = 50
MAX_WORKBOOK_CELLS = 200_000
DEFAULT_SHEET_ROWS = 100
DEFAULT_SHEET_COLS = 26


def _argb_color(value: Any) -> str | None:
    """openpyxl 的颜色是 ARGB；只接受不透明的 #RRGGBB。"""
    raw = str(value or "").strip().lstrip("#")
    if len(raw) == 8:
        if raw[:2].upper() == "00":
            return None
        raw = raw[2:]
    if len(raw) != 6 or not re.fullmatch(r"[0-9a-fA-F]{6}", raw):
        return None
    if raw.upper() == "000000":
        return None
    return f"#{raw.upper()}"


def _number_format(value: Any) -> str | None:
    fmt = str(value or "").strip().lower()
    if not fmt or fmt == "general":
        return None
    if "%" in fmt:
        return "percent"
    if any(token in fmt for token in ("yy", "mm", "dd", "hh", "ss", "年", "月", "日")):
        return "date"
    if any(token in fmt for token in ("¥", "￥", "$", "€", "£", "currency")):
        return "currency"
    if fmt == "@":
        return "text"
    if any(token in fmt for token in ("0", "#")):
        return "number"
    return None


def _cell_style(cell: Any, *, bold: bool | None = None) -> dict:
    """把 openpyxl 单元格样式压缩成前端渲染需要的少量字段。"""
    style: dict[str, Any] = {}
    font = getattr(cell, "font", None)
    if font is not None:
        if bold is not None:
            if bold:
                style["bold"] = True
        elif font.bold:
            style["bold"] = True
        if font.italic:
            style["italic"] = True
        if font.underline:
            style["underline"] = True
        color = _argb_color(getattr(font, "color", None) and font.color.rgb)
        if color:
            style["color"] = color
    fill = getattr(cell, "fill", None)
    if fill is not None and getattr(fill, "fill_type", None) == "solid":
        background = _argb_color(fill.fgColor.rgb if fill.fgColor else None)
        if background:
            style["bg"] = background
    alignment = getattr(cell, "alignment", None)
    if alignment is not None and alignment.horizontal in ("left", "center", "right"):
        style["align"] = alignment.horizontal
    fmt = _number_format(getattr(cell, "number_format", None))
    if fmt:
        style["fmt"] = fmt
    return style


def _sheet_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, datetime):
        # 去掉无意义的 00:00:00，保持与 Excel 一致的日期观感。
        return value.isoformat(sep=" ") if (value.hour or value.minute or value.second) else value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, (int, float)):
        return value
    text = str(value)
    return text[:MAX_CELL_CHARS]


def _date_style(value: Any) -> dict:
    if isinstance(value, (date, datetime)):
        return {"fmt": "date"}
    return {}


def _grid_workbook(
    names: list[str], sheets: dict[str, list[list[Any]]], *, has_header: bool
) -> tuple[dict, bool]:
    """CSV / .xls 的通用回退：无公式与样式，仅保留单元格文本。"""
    sheet_list: list[dict] = []
    truncated = False
    for name in names[:MAX_SHEETS]:
        grid = sheets.get(name, [])
        rows = max(DEFAULT_SHEET_ROWS, min(len(grid) + 10, MAX_TABLE_ROWS))
        cols = max(
            DEFAULT_SHEET_COLS,
            min(max((len(row) for row in grid), default=0), MAX_TABLE_COLUMNS),
        )
        cells: dict[str, dict] = {}
        for row_index, row in enumerate(grid[:MAX_TABLE_ROWS]):
            for col_index, raw in enumerate(row[:MAX_TABLE_COLUMNS]):
                if raw is None or (isinstance(raw, str) and not raw.strip()):
                    continue
                value = _sheet_value(raw)
                style = _date_style(raw)
                if has_header and row_index == 0:
                    style["bold"] = True
                cell: dict[str, Any] = {"v": value}
                if style:
                    cell["s"] = style
                cells[_address(row_index, col_index)] = cell
        sheet_list.append(
            {
                "id": uuid.uuid4().hex,
                "name": name[:80] or f"Sheet{len(sheet_list) + 1}",
                "rows": rows,
                "cols": cols,
                "cells": cells,
                "colWidths": {},
                "rowHeights": {},
                "merges": [],
            }
        )
    if not sheet_list:
        raise TableImportError("表格文件没有可读取的工作表", 422)
    return {
        "version": 1,
        "active": sheet_list[0]["id"],
        "sheets": sheet_list,
    }, truncated


def _address(row: int, col: int) -> str:
    """0 基行列 → A1 地址。"""
    label = ""
    value = col + 1
    while value > 0:
        value, rem = divmod(value - 1, 26)
        label = chr(65 + rem) + label
    return f"{label}{row + 1}"


def _workbook_from_xlsx(content: bytes) -> tuple[dict, bool]:
    _guard_xlsx(content)
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - 依赖缺失
        raise TableImportError("服务端缺少 Excel 解析组件 openpyxl", 500) from exc

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=False, data_only=False)
        cached_book = load_workbook(io.BytesIO(content), read_only=False, data_only=True)
    except Exception as exc:  # noqa: BLE001 - 第三方解析异常统一转 415
        raise TableImportError("Excel 文件解析失败，请确认文件未加密", 415) from exc

    try:
        from openpyxl.utils import column_index_from_string, coordinate_to_tuple
    except ImportError as exc:  # pragma: no cover - 依赖缺失
        workbook.close()
        cached_book.close()
        raise TableImportError("服务端缺少 Excel 解析组件 openpyxl", 500) from exc

    sheet_list: list[dict] = []
    truncated = False
    total_cells = 0
    try:
        for worksheet in workbook.worksheets[:MAX_SHEETS]:
            cached_ws = cached_book[worksheet.title]
            max_row = min(worksheet.max_row or 0, MAX_TABLE_ROWS)
            max_col = min(worksheet.max_column or 0, MAX_TABLE_COLUMNS)
            if worksheet.max_row and worksheet.max_row > MAX_TABLE_ROWS:
                truncated = True
            if worksheet.max_column and worksheet.max_column > MAX_TABLE_COLUMNS:
                truncated = True
            rows = max(DEFAULT_SHEET_ROWS, min(max_row + 10, MAX_TABLE_ROWS))
            cols = max(DEFAULT_SHEET_COLS, min(max_col, MAX_TABLE_COLUMNS))
            cells: dict[str, dict] = {}
            for row_index in range(max_row):
                for col_index in range(max_col):
                    cell = worksheet.cell(row=row_index + 1, column=col_index + 1)
                    raw = cell.value
                    if raw is None and not cell.has_style:
                        continue
                    if total_cells >= MAX_WORKBOOK_CELLS:
                        truncated = True
                        break
                    entry: dict[str, Any] = {}
                    if isinstance(raw, str) and raw.startswith("="):
                        entry["f"] = raw[:MAX_CELL_CHARS]
                        cached = cached_ws.cell(
                            row=row_index + 1, column=col_index + 1
                        ).value
                        cached_value = _sheet_value(cached)
                        if cached_value is not None:
                            entry["cached"] = cached_value
                    else:
                        value = _sheet_value(raw)
                        if value is not None:
                            entry["v"] = value
                    style = _cell_style(cell)
                    style.update(_date_style(raw))
                    if style:
                        entry["s"] = style
                    if not entry:
                        continue
                    cells[_address(row_index, col_index)] = entry
                    total_cells += 1
                if total_cells >= MAX_WORKBOOK_CELLS:
                    break

            col_widths: dict[str, int] = {}
            for letter, dimension in worksheet.column_dimensions.items():
                width = getattr(dimension, "width", None)
                if not width:
                    continue
                try:
                    index = column_index_from_string(letter) - 1
                except Exception:  # noqa: BLE001 - 异常维度名直接跳过
                    continue
                if 0 <= index < cols:
                    col_widths[str(index)] = max(48, min(420, int(width * 7.5 + 5)))

            row_heights: dict[str, int] = {}
            for index, dimension in worksheet.row_dimensions.items():
                height = getattr(dimension, "height", None)
                if not height:
                    continue
                if 1 <= index <= rows:
                    row_heights[str(index - 1)] = max(20, min(200, int(height * 4 / 3)))

            merges: list[str] = []
            for merged in worksheet.merged_cells.ranges:
                if merged.min_row > MAX_TABLE_ROWS or merged.min_col > MAX_TABLE_COLUMNS:
                    truncated = True
                    continue
                merges.append(str(merged))

            # Excel 的冻结窗格保存为分隔线左上角单元格，如 "B2" 表示冻结
            # 第 1 行和第 A 列；换算成前端的冻结数量。
            frozen_rows = 0
            frozen_cols = 0
            freeze_panes = worksheet.freeze_panes
            if freeze_panes:
                try:
                    freeze_row, freeze_col = coordinate_to_tuple(str(freeze_panes))
                except (TypeError, ValueError):
                    freeze_row, freeze_col = 1, 1
                frozen_rows = max(0, min(freeze_row - 1, rows - 1))
                frozen_cols = max(0, min(freeze_col - 1, cols - 1))

            sheet_list.append(
                {
                    "id": uuid.uuid4().hex,
                    "name": worksheet.title[:80] or f"Sheet{len(sheet_list) + 1}",
                    "rows": rows,
                    "cols": cols,
                    "cells": cells,
                    "colWidths": col_widths,
                    "rowHeights": row_heights,
                    "merges": merges,
                    "frozenRows": frozen_rows,
                    "frozenCols": frozen_cols,
                }
            )
    finally:
        workbook.close()
        cached_book.close()

    if not sheet_list:
        raise TableImportError("Excel 中没有可读取的工作表", 422)
    return {
        "version": 1,
        "active": sheet_list[0]["id"],
        "sheets": sheet_list,
    }, truncated


def parse_sheet_workbook(
    content: bytes, filename: str, *, has_header: bool = False
) -> tuple[dict, bool]:
    """把 CSV / Excel 解析成表格笔记的工作簿 JSON。"""
    if not content:
        raise TableImportError("表格文件为空", 400)
    if len(content) > MAX_TABLE_BYTES:
        raise TableImportError("表格文件超过 20MB 上限", 413)
    scan_status, scan_detail = scan_bytes(content)
    if scan_status in ("infected", "error"):
        raise TableImportError(f"文件安全扫描未通过：{scan_detail}", 422)

    ext = Path(filename).suffix.lower()
    if ext == ".xlsx":
        return _workbook_from_xlsx(content)
    if ext in (".csv", ".xls"):
        names, sheets = _read_workbook(content, filename)
        return _grid_workbook(names, sheets, has_header=has_header)
    raise TableImportError(
        f"不支持的表格类型 {ext or '（无扩展名）'}，允许：.csv、.xlsx、.xls", 415
    )


def parse_table(
    content: bytes,
    filename: str,
    *,
    sheet: str | None = None,
    has_header: bool = True,
) -> ParsedTable:
    """解析表格并推断字段类型；保留所有工作表供前端切换。"""
    if not content:
        raise TableImportError("表格文件为空", 400)
    if len(content) > MAX_TABLE_BYTES:
        raise TableImportError("表格文件超过 20MB 上限", 413)

    scan_status, scan_detail = scan_bytes(content)
    if scan_status in ("infected", "error"):
        raise TableImportError(f"文件安全扫描未通过：{scan_detail}", 422)

    names, sheets = _read_workbook(content, filename)
    non_empty = [
        name
        for name in names
        if any(any(not _is_empty(cell) for cell in row) for row in sheets.get(name, []))
    ]
    candidates = non_empty or names
    chosen = sheet if sheet in sheets else (candidates[0] if candidates else "")
    if not chosen:
        raise TableImportError("Excel 中没有可读取的工作表", 422)

    grid = [row for row in sheets.get(chosen, []) if any(not _is_empty(cell) for cell in row)]
    if not grid:
        raise TableImportError(f"工作表「{chosen}」没有数据", 422)
    truncated = len(grid) > MAX_TABLE_ROWS + (1 if has_header else 0)
    if truncated:
        grid = grid[: MAX_TABLE_ROWS + (1 if has_header else 0)]

    width = min(max(len(row) for row in grid), MAX_TABLE_COLUMNS)
    if width == 0:
        raise TableImportError("表格没有可导入的列", 422)
    grid = [(row + [""] * width)[:width] for row in grid]

    if has_header:
        headers = _normalize_headers(grid[0], width)
        data_rows = grid[1:]
    else:
        headers = _normalize_headers([], width)
        data_rows = grid

    fields: list[dict] = []
    for index, header in enumerate(headers):
        field_id = f"c{index + 1}"
        values = [row[index] for row in data_rows]
        fields.append(_infer(field_id, header, values))

    rows: list[dict] = []
    for row in data_rows[:MAX_TABLE_ROWS]:
        if all(_is_empty(cell) for cell in row):
            continue
        rows.append(
            {
                field["id"]: _typed_value(field, row[index])
                for index, field in enumerate(fields)
            }
        )

    sheet_infos = [
        SheetInfo(
            name=name,
            row_count=max(0, len([r for r in sheets.get(name, []) if any(not _is_empty(c) for c in r)]) - (1 if has_header else 0)),
            column_count=min(
                max((len(r) for r in sheets.get(name, [])), default=0),
                MAX_TABLE_COLUMNS,
            ),
        )
        for name in names
    ]
    return ParsedTable(
        sheets=sheet_infos,
        sheet=chosen,
        fields=fields,
        rows=rows,
        truncated=truncated or len(rows) > MAX_TABLE_ROWS,
    )


def preview_payload(table: ParsedTable) -> dict:
    return {
        "sheets": [
            {"name": item.name, "row_count": item.row_count, "column_count": item.column_count}
            for item in table.sheets
        ],
        "sheet": table.sheet,
        "fields": table.fields,
        "sample_rows": table.rows[:MAX_SAMPLE_ROWS],
        "row_count": len(table.rows),
        "truncated": table.truncated,
    }


__all__ = [
    "TABLE_EXT",
    "MAX_TABLE_BYTES",
    "TableImportError",
    "ParsedTable",
    "parse_sheet_workbook",
    "parse_table",
    "preview_payload",
]
