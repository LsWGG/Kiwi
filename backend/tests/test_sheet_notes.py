"""表格笔记：创建、Excel / CSV 导入以及工作簿结构保留。"""

from __future__ import annotations

import io
import uuid

from openpyxl import Workbook
from openpyxl.styles import PatternFill


def _xlsx_bytes() -> bytes:
    book = Workbook()
    first = book.active
    first.title = "季度数据"
    first.append(["项目", "Q1", "Q2", "合计"])
    first.append(["收入", 120, 180, "=B2+C2"])
    first.append(["成本", 80, 90, "=B3+C3"])
    first["A1"].fill = PatternFill("solid", fgColor="FFDBEAFE")
    first.merge_cells("A1:A2")
    second = book.create_sheet("说明")
    second.append(["备注"])
    second.append(["多工作表导入"])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


async def test_create_sheet_note(client, kb):
    workbook = {
        "version": 1,
        "active": "s1",
        "sheets": [
            {
                "id": "s1",
                "name": "Sheet1",
                "rows": 20,
                "cols": 8,
                "cells": {"A1": {"v": "名称"}, "B2": {"f": "=1+1"}},
                "colWidths": {},
                "rowHeights": {},
                "merges": [],
            }
        ],
    }
    resp = await client.post(
        f"/api/notes?kb_id={kb['id']}",
        json={
            "title": f"sheet-{uuid.uuid4().hex[:6]}",
            "kind": "sheet",
            "props_json": {"sheet": workbook},
        },
    )
    assert resp.status_code == 201, resp.text
    note = resp.json()
    assert note["kind"] == "sheet"
    assert note["props_json"]["sheet"]["sheets"][0]["cells"]["B2"]["f"] == "=1+1"


async def test_import_table_as_sheet_keeps_workbook(client, kb):
    resp = await client.post(
        f"/api/notes/import-table?kb_id={kb['id']}",
        data={"target": "sheet", "has_header": "true"},
        files={
            "file": (
                "quarter.xlsx",
                _xlsx_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert resp.status_code == 201, resp.text
    data = resp.json()
    note = data["note"]
    assert note["kind"] == "sheet"
    sheet = note["props_json"]["sheet"]
    assert len(sheet["sheets"]) == 2
    assert sheet["sheets"][0]["name"] == "季度数据"
    cells = sheet["sheets"][0]["cells"]
    assert cells["D2"]["f"] == "=B2+C2"
    assert cells["A1"]["s"]["bg"] == "#DBEAFE"
    assert sheet["sheets"][0]["merges"] == ["A1:A2"]
    assert note["props_json"]["table_import"]["sheet_count"] == 2


async def test_import_table_as_sheet_csv(client, kb):
    csv = "名称,数量\n苹果,3\n香蕉,5\n".encode()
    resp = await client.post(
        f"/api/notes/import-table?kb_id={kb['id']}",
        data={"target": "sheet"},
        files={"file": ("fruit.csv", csv, "text/csv")},
    )
    assert resp.status_code == 201, resp.text
    note = resp.json()["note"]
    assert note["kind"] == "sheet"
    sheet = note["props_json"]["sheet"]["sheets"][0]
    assert sheet["cells"]["A1"]["v"] == "名称"
    assert sheet["cells"]["B2"]["v"] == "3"


async def test_import_table_rejects_unknown_target(client, kb):
    resp = await client.post(
        f"/api/notes/import-table?kb_id={kb['id']}",
        data={"target": "pdf"},
        files={"file": ("fruit.csv", b"a,b\n1,2\n", "text/csv")},
    )
    assert resp.status_code == 400
